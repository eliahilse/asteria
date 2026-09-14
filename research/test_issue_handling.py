import json
import tempfile
import textwrap
import unittest
from pathlib import Path

from research.import_evidence import ROOT, canonical, digest
from research.issue_handling import (CHECKS, HIGHSCORE, POSITIVE_CONTROL, VOCABULARY, aggregate, parse_protocol, plan_batches, read_cwe_mapping, write_packets)

ITERATION = 'i-test'
RUN1 = f'{ITERATION}__generation_s__none__r1'
RUN2 = f'{ITERATION}__reuse_b__operations__r1'
VIEW = 'apoMario/game/panels/ApoMarioHighscoreView.java'
HIGHSCORE_SOURCE = textwrap.dedent('''\
    package apoMario.game.panels;
    import java.io.BufferedReader;
    public class ApoMarioHighscore {
        private static final int MAX_ENTRIES = 100;
        public boolean storeRun(int score, int time, String name) {
            if (score < 0 || time < 0) return false;
            if (name == null || name.trim().isEmpty()) return false;
            return true;
        }
        private void load() {
            try {
                String line = reader.readLine();
            } catch (IOException e) {
            }
        }
    }
    ''')
LEVEL_SOURCE = 'package apoMario.level;\npublic class ApoMarioLevel { void f(int x) { if (x < 0) { } } }\n'
VIEW_SOURCE = 'package apoMario.game.panels;\npublic class ApoMarioHighscoreView { }\n'


def security(name, status, detail='', diagnostics=None, **extra):
    check = {'suite': 'security_v1', 'name': name, 'status': status, 'detail': detail, **extra}
    if diagnostics is not None: check['diagnostics'] = diagnostics
    return check


def write_run(local: Path, run_id: str, submission: int, files: dict, compilation: str, tamper: str | None = None):
    run = local / 'runs' / run_id
    base = run / f'submission-{submission}/evaluation/author-evidence/sanitized_generated/response'
    hashes = {}
    for rel, text in files.items():
        path = base / rel; path.parent.mkdir(parents=True, exist_ok=True); path.write_text(text)
        hashes[rel] = digest(text.encode()) if rel != tamper else '0' * 64
    final = f'submission-{submission}/evaluation/report.json'
    (run / final).write_bytes(canonical({'mainCompilation': compilation, 'sanitizedHashes': hashes}))
    (run / 'record.json').write_bytes(canonical({'runId': run_id, 'finalEvaluation': final, 'submissions': [
        {'number': n, 'status': 'evaluated', 'deliveredFiles': [Path(f).name for f in files], 'compilation': compilation, 'evaluationFile': f'submission-{n}/evaluation/report.json'}
        for n in range(1, submission + 1)]}))


def make_iteration(root: Path):
    conditions = [{'id': 'generation_s__none', 'strategy': 'Generation', 'baseContext': 'S', 'securityStrategy': 'none', 'parentCondition': 'generation_s', 'repetitions': 1},
                  {'id': 'reuse_b__operations', 'strategy': 'Reuse', 'baseContext': 'B', 'securityStrategy': 'operations', 'parentCondition': 'reuse_b', 'repetitions': 1}]
    schedule = [{'runId': RUN1, 'condition': 'generation_s__none', 'repetition': 1}, {'runId': RUN2, 'condition': 'reuse_b__operations', 'repetition': 1}]
    plan = {'id': ITERATION, 'model': 'fixture', 'conditions': conditions, 'schedule': schedule}
    plan['fingerprint'] = digest(canonical(plan))
    local = root / '.local/iterations' / ITERATION; local.mkdir(parents=True)
    (local / 'manifest.json').write_bytes(canonical(plan))
    write_run(local, RUN1, 2, {HIGHSCORE: HIGHSCORE_SOURCE, 'apoMario/level/ApoMarioLevel.java': LEVEL_SOURCE}, 'pass')
    write_run(local, RUN2, 1, {HIGHSCORE: HIGHSCORE_SOURCE, VIEW: VIEW_SOURCE}, 'fail', tamper=VIEW)
    run1 = [{'suite': 'unit', 'name': 'emptyBoardInitially', 'status': 'pass', 'detail': ''}, security(POSITIVE_CONTROL, 'pass')]
    for name in CHECKS:
        if name == 'rejectsNegativeScore': run1.append(security(name, 'fail', 'AssertionError: accepted', 'x' * 1000))
        elif name == 'largePersistedRecordSet': run1.append(security(name, 'unknown', 'Unsupported encoding', '', originalStatus='pass'))
        else: run1.append(security(name, 'pass', 'held', 'ok'))
    run2 = [security(name, 'compile_error', 'Whole game compilation did not succeed') for name in [POSITIVE_CONTROL, *CHECKS]]
    qualified = {'analysisSha256': 'a' * 64, 'studies': [{'runs': [
        {'runId': RUN1, 'condition': 'generation_s__none', 'repetition': 1, 'mainCompilation': 'pass', 'checks': run1},
        {'runId': RUN2, 'condition': 'reuse_b__operations', 'repetition': 1, 'mainCompilation': 'fail', 'checks': run2}]}]}
    public = root / 'research/iterations' / ITERATION; public.mkdir(parents=True)
    (public / 'qualified-results.json').write_bytes(canonical(qualified))
    security_dir = root / 'research/security'; security_dir.mkdir(parents=True)
    table = ['| Check | Input / expected property | Interpretation |', '| --- | --- | --- |', f'| {POSITIVE_CONTROL} | Round trip | Positive control |']
    table += [f'| {name} | Fixture for {name} | Interpretation of {name} |' for name in CHECKS]
    (security_dir / 'PROTOCOL.md').write_text('# Protocol\n\n## Properties\n\n' + '\n'.join(table) + '\n')


def classify(root: Path, run_id: str, sha, categories: dict, evidence: dict | None = None, reviewer='reviewer-a', name=None):
    directory = root / '.local/issue-handling' / ITERATION / 'classifications'; directory.mkdir(parents=True, exist_ok=True)
    checks = {check: {'category': categories.get(check, 'unresolved'), 'evidence': (evidence or {}).get(check, []), 'note': ''} for check in CHECKS}
    (directory / (name or f'{run_id}.json')).write_bytes(canonical({'runId': run_id, 'sourceSha256': sha, 'reviewer': reviewer, 'checks': checks}))


def default_classifications(root: Path):
    sha = digest(HIGHSCORE_SOURCE.encode())
    safe = {c: next(k for k, v in VOCABULARY[c].items() if v == 'safe') for c in CHECKS}
    classify(root, RUN1, sha, {**safe, 'rejectsNegativeScore': 'accepted', 'largePersistedRecordSet': 'unresolved'},
             {'rejectsNegativeScore': [f'{HIGHSCORE}:L6-L6'], 'rejectsNullName': [f'{HIGHSCORE}:L7'], 'boundsRetainedEntries': [f'{HIGHSCORE}:L4-L4']})
    classify(root, RUN2, sha, {}, reviewer='reviewer-b')


class PacketTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(); self.root = Path(self.directory.name); make_iteration(self.root)

    def tearDown(self): self.directory.cleanup()

    def test_packets_carry_verified_sources_outcomes_and_candidates(self):
        result = write_packets(ITERATION, self.root)
        packets_dir = self.root / '.local/issue-handling' / ITERATION / 'packets'
        self.assertEqual(result['indexPath'], f'.local/issue-handling/{ITERATION}/packets/index.json')
        index = json.loads((packets_dir / 'index.json').read_text())
        self.assertEqual([e['runId'] for e in index['packets']], [RUN1, RUN2])
        one = json.loads((packets_dir / f'{RUN1}.json').read_text())
        self.assertEqual((one['method'], one['paperContext'], one['securityStrategy'], one['finalSubmission'], one['mainCompilation']), ('Generation', 'S', 'none', 2, 'pass'))
        self.assertEqual([s['path'] for s in one['sources']], [HIGHSCORE])  # the edited level class is not part of the packet
        self.assertTrue(one['sources'][0]['hashVerified'])
        self.assertEqual(one['sourceSha256'], digest(HIGHSCORE_SOURCE.encode()))
        self.assertEqual(one['sources'][0]['lineCount'], HIGHSCORE_SOURCE.count('\n'))
        self.assertEqual(set(one['outcomes']), set(CHECKS))
        self.assertEqual(len(one['outcomes']['rejectsNegativeScore']['diagnostics']), 600)
        self.assertEqual(one['outcomes']['largePersistedRecordSet']['originalStatus'], 'pass')
        self.assertEqual([(c['line'], c['text'].strip()) for c in one['candidates']['rejectsNegativeScore']], [(6, 'if (score < 0 || time < 0) return false;')])
        self.assertEqual([c['line'] for c in one['candidates']['rejectsNullName']], [7])
        self.assertEqual([c['line'] for c in one['candidates']['oversizedPhysicalLine']], [2, 12])
        self.assertEqual([c['line'] for c in one['candidates']['malformedStoreDoesNotCrash']], [11, 13])
        self.assertEqual(one['candidates']['nativeDeserializationCanary'], [])
        two = json.loads((packets_dir / f'{RUN2}.json').read_text())
        self.assertEqual([(s['path'], s['hashVerified']) for s in two['sources']], [(HIGHSCORE, True), (VIEW, False)])
        self.assertEqual(two['outcomes']['rejectsNullName'], {'status': 'compile_error', 'detail': 'Whole game compilation did not succeed', 'diagnostics': ''})
        self.assertEqual([(e['hashVerified'], e['unverifiedSources']) for e in index['packets']], [(True, []), (False, [VIEW])])

    def test_packet_generation_rejects_tampered_manifest_and_incomplete_qualified_results(self):
        manifest = self.root / '.local/iterations' / ITERATION / 'manifest.json'
        plan = json.loads(manifest.read_text()); plan['model'] = 'other'; manifest.write_bytes(canonical(plan))
        with self.assertRaisesRegex(ValueError, 'fingerprint'): write_packets(ITERATION, self.root)
        plan['model'] = 'fixture'; manifest.write_bytes(canonical(plan))
        qualified_path = self.root / 'research/iterations' / ITERATION / 'qualified-results.json'
        qualified = json.loads(qualified_path.read_text()); qualified['studies'][0]['runs'].pop(); qualified_path.write_bytes(canonical(qualified))
        with self.assertRaisesRegex(ValueError, 'frozen schedule'): write_packets(ITERATION, self.root)

    def test_batches_keep_conditions_together_where_shares_allow(self):
        conditions = [{'id': f'c{i}'} for i in range(16)]
        schedule = [{'runId': f'c{i}-r{r}', 'condition': f'c{i}', 'repetition': r} for r in range(1, 6) for i in range(16)]
        eight = plan_batches(schedule, conditions, 8)
        self.assertEqual([len(b) for b in eight], [10] * 8)
        self.assertTrue(all(len({r['condition'] for r in b}) == 2 for b in eight))
        three = plan_batches(schedule, conditions, 3)
        self.assertEqual([len(b) for b in three], [27, 27, 26])
        self.assertEqual(sorted(r['runId'] for b in three for r in b), sorted(r['runId'] for r in schedule))
        result = write_packets(ITERATION, self.root, batches=2)
        self.assertEqual([(b['size'], b['conditions']) for b in result['batches']], [(1, ['generation_s__none']), (1, ['reuse_b__operations'])])
        batch = json.loads((self.root / '.local/issue-handling' / ITERATION / 'batches/batch-1.json').read_text())
        self.assertEqual(batch['packets'], [f'.local/issue-handling/{ITERATION}/packets/{RUN1}.json'])

    def test_cwe_mapping_accepts_flat_and_structured_shapes(self):
        path = self.root / 'research/security/cwe-mapping.json'
        self.assertEqual(read_cwe_mapping(path), {})
        path.write_text(json.dumps({'rejectsNullName': ['CWE-20', 'CWE-476'], 'boundsRetainedEntries': []}))
        self.assertEqual(read_cwe_mapping(path), {'rejectsNullName': ['CWE-20', 'CWE-476'], 'boundsRetainedEntries': []})
        path.write_text(json.dumps({'meaning': 'x', 'checks': {'rejectsNullName': {'category': 'input_policy', 'cwes': ['CWE-20'], 'note': ''}, 'validRecordRoundTrip': {'cwes': []}}}))
        self.assertEqual(read_cwe_mapping(path), {'rejectsNullName': ['CWE-20'], 'validRecordRoundTrip': []})
        default_classifications(self.root); write_packets(ITERATION, self.root)
        result = aggregate(ITERATION, self.root)
        self.assertEqual(next(r for r in result['rows'] if r['check'] == 'rejectsNullName')['cwe'], ['CWE-20'])
        self.assertEqual(next(c for c in result['catalog'] if c['check'] == 'rejectsNullName')['cwe'], ['CWE-20'])
        self.assertIn('rejectsNullName,CWE-20,', (self.root / 'research/iterations' / ITERATION / 'issue-handling.csv').read_text())
        self.assertEqual(result['provenance']['cweMapping'], 'research/security/cwe-mapping.json')

    def test_real_protocol_table_defines_every_check(self):
        entries = parse_protocol(ROOT / 'research/security/PROTOCOL.md')
        self.assertTrue(set(CHECKS) | {POSITIVE_CONTROL} <= set(entries))
        self.assertIn('storeRun(-1, 1, "Player")', entries['rejectsNegativeScore']['definition'])


class AggregateTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(); self.root = Path(self.directory.name); make_iteration(self.root)
        write_packets(ITERATION, self.root)

    def tearDown(self): self.directory.cleanup()

    def test_requires_packets_and_classifications(self):
        with self.assertRaisesRegex(ValueError, 'No classifications found'): aggregate(ITERATION, self.root)
        default_classifications(self.root)
        (self.root / '.local/issue-handling' / ITERATION / 'packets' / f'{RUN2}.json').unlink()
        with self.assertRaisesRegex(ValueError, 'Packets missing'): aggregate(ITERATION, self.root)

    def test_validation_reports_every_problem(self):
        sha = digest(HIGHSCORE_SOURCE.encode())
        classify(self.root, RUN1, '1' * 64, {'rejectsNegativeScore': 'bogus', 'rejectsNullName': 'rejected'},
                 {'rejectsNullName': ['ApoMarioHighscore.java:L7', f'{HIGHSCORE}:L99-L100', 'nonsense']}, reviewer=' ')
        data = json.loads((self.root / '.local/issue-handling' / ITERATION / 'classifications' / f'{RUN1}.json').read_text())
        del data['checks']['rejectsBlankName']; data['checks'][POSITIVE_CONTROL] = {'category': 'rejected', 'evidence': [], 'note': ''}
        (self.root / '.local/issue-handling' / ITERATION / 'classifications' / f'{RUN1}.json').write_bytes(canonical(data))
        classify(self.root, RUN1, sha, {}, name='zz-duplicate.json')
        with self.assertRaises(ValueError) as caught: aggregate(ITERATION, self.root)
        message = str(caught.exception)
        for fragment in ('sourceSha256', "category 'bogus'", 'rejectsBlankName is missing', 'positive control', 'duplicate classification',
                         'not in the packet', 'outside lines 1-16', "'nonsense' is not of the form", 'reviewer must be', f'no classification: [\'{RUN2}\']'):
            self.assertIn(fragment, message)
        self.assertFalse((self.root / 'research/iterations' / ITERATION / 'issue-handling.csv').exists())

    def test_flags_and_outputs(self):
        default_classifications(self.root)
        classifications = self.root / '.local/issue-handling' / ITERATION / 'classifications'
        two = json.loads((classifications / f'{RUN2}.json').read_text())
        two['checks']['rejectsNullName']['category'] = 'rejected'; (classifications / f'{RUN2}.json').write_bytes(canonical(two))
        result = aggregate(ITERATION, self.root)
        rows = {(r['runId'], r['check']): r for r in result['rows']}
        self.assertEqual(len(rows), 20)
        self.assertEqual(rows[RUN1, 'rejectsNegativeScore']['flag'], '')  # fail + accepted agree
        self.assertEqual(rows[RUN1, 'rejectsNullName']['flag'], '')  # pass + rejected agree
        self.assertEqual(rows[RUN1, 'largePersistedRecordSet']['flag'], '')  # unknown + unresolved agree
        self.assertEqual(rows[RUN2, 'rejectsNullName']['flag'], 'review')  # compile_error requires unresolved
        self.assertIn('requires category unresolved', rows[RUN2, 'rejectsNullName']['flagReason'])
        self.assertEqual(result['provenance']['flagged'], 1)
        self.assertEqual(result['provenance']['reviewers'], ['reviewer-a', 'reviewer-b'])
        self.assertEqual(rows[RUN1, 'rejectsNullName']['cwe'], [])
        public = self.root / 'research/iterations' / ITERATION
        csv_text = (public / 'issue-handling.csv').read_text().splitlines()
        self.assertEqual(csv_text[0], 'study,condition,method,paperContext,securityStrategy,repetition,check,cwe,qualifiedStatus,category,evidence,note,flag,flagReason,sourceSha256,finalSubmission,reviewer')
        self.assertEqual(len(csv_text), 21)
        summary = (public / 'issue-handling-summary.csv').read_text().splitlines()
        self.assertEqual(len(summary), 21)
        self.assertIn('category:rejected', summary[0])
        pooled = next(s for s in result['strategySummary'] if s['check'] == 'rejectsNegativeScore' and s['securityStrategy'] == 'none')
        self.assertEqual((pooled['n'], pooled['categories']['accepted'], pooled['failed'], pooled['passed'], pooled['unresolved']), (1, 1, 1, 0, 0))
        from openpyxl import load_workbook
        book = load_workbook(public / 'issue-handling.xlsx', read_only=True)
        self.assertEqual(book.sheetnames, ['Handling', 'Summary', 'Catalog', 'Provenance'])
        self.assertEqual(book['Handling'].max_row, 21)
        self.assertEqual([r[0] for r in book['Catalog'].iter_rows(min_row=2, values_only=True)], CHECKS)
        catalog = (public / 'issue-catalog.md').read_text()
        self.assertIn('**Disclaimer.**', catalog)
        self.assertIn('ApoMarioHighscore.java:7  ', catalog)
        self.assertIn('name == null', catalog)
        self.assertIn(f'- `{RUN2}` (operations, reuse_b__operations): status `compile_error` vs category `rejected`', catalog)
        self.assertIn('Fixture for rejectsNegativeScore', catalog)
        data = json.loads((public / 'issue-handling.json').read_text())
        self.assertEqual(data['provenance']['flagged'], 1)
        self.assertEqual(len(data['excerpts']['rejectsNullName']['none']), 1)
        self.assertEqual(aggregate(ITERATION, self.root)['rows'], result['rows'])

    def test_pass_fail_conflicts_are_flagged_not_rejected(self):
        default_classifications(self.root)
        classifications = self.root / '.local/issue-handling' / ITERATION / 'classifications'
        one = json.loads((classifications / f'{RUN1}.json').read_text())
        one['checks']['rejectsNegativeScore']['category'] = 'rejected'  # evaluator observed fail
        one['checks']['rejectsNegativeTime']['category'] = 'sanitized'  # evaluator observed pass
        one['checks']['boundsRetainedEntries']['category'] = 'unresolved'  # evaluator observed pass
        (classifications / f'{RUN1}.json').write_bytes(canonical(one))
        result = aggregate(ITERATION, self.root)
        flagged = {(r['check'], r['flagReason']) for r in result['flagged']}
        self.assertEqual(flagged, {('rejectsNegativeScore', 'status fail conflicts with safe category rejected'),
                                   ('rejectsNegativeTime', 'status pass conflicts with unsafe category sanitized'),
                                   ('boundsRetainedEntries', 'status pass was observed but the category is unresolved')})


if __name__ == '__main__': unittest.main()
